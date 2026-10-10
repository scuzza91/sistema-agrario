# -*- coding: utf-8 -*-
"""Genera la documentacion de la segunda entrega (MongoDB) a partir de la plantilla de la catedra.

Uso (con MongoDB levantado, `npm run seed` ejecutado y el servidor en http://localhost:3000):
    python docs/generar_entrega.py
"""

import json
import os
import urllib.error
import urllib.request

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = "http://localhost:3000"
ROOT = os.path.dirname(os.path.abspath(__file__))
PROYECTO = os.path.dirname(ROOT)
NOMBRE = "DSWB_9_MonitoreoLotes_1C26"
SALIDA = os.path.join(os.path.dirname(PROYECTO), NOMBRE + "_Parte2.docx")
LOGO = os.path.join(ROOT, "ifts29.png")
CAPTURAS = [
    ("Pruebas de la API (resultado de este script, ver docs/pruebas.html)", os.path.join(ROOT, "captura-pruebas.png")),
    ("Vista Pug de lotes: los datos vienen de MongoDB con populate", os.path.join(ROOT, "vista-lotes-pug.png")),
    ("Vista Pug de observaciones", os.path.join(ROOT, "vista-observaciones-pug.png")),
    ("Error de validacion mostrado desde un formulario", os.path.join(ROOT, "captura-error-validacion.png")),
]

DRIVE = "https://drive.google.com/drive/folders/1xIs78dB4CbTy_byZE-5crOYv3pvAp3JZ?usp=sharing"
GITHUB = "https://github.com/scuzza91/sistema-agrario"
VIDEO = "COMPLETAR: link al video (Meet / Zoom / YouTube) con permiso de visualizacion"

AZUL = (0x1F, 0x4E, 0x79)


def pedir(metodo, ruta, cuerpo=None):
    data = None
    headers = {"Accept": "application/json"}
    if cuerpo is not None:
        data = json.dumps(cuerpo).encode("utf-8")
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(BASE + ruta, data=data, headers=headers, method=metodo)
    try:
        with urllib.request.urlopen(req) as resp:
            return resp.status, json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as err:
        return err.code, json.loads(err.read().decode("utf-8"))


def set_run_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def titulo(doc, text, size=16):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), "Calibri", size, True, AZUL)


def parrafo(doc, text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    set_run_font(p.add_run(text), "Calibri", 11, bold)
    return p


def vineta(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_run_font(p.add_run(text), "Calibri", 11)


def codigo(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    set_run_font(p.add_run(text), "Consolas", 9, False, (0x33, 0x33, 0x33))


def sombrear(cell, hex_color):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def tabla(doc, encabezados, filas, anchos=None):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.autofit = anchos is None
    for i, h in enumerate(encabezados):
        c = t.rows[0].cells[i]
        c.text = ""
        set_run_font(c.paragraphs[0].add_run(h), "Calibri", 10, True, (0xFF, 0xFF, 0xFF))
        sombrear(c, "1F4E79")
    for fila in filas:
        celdas = t.add_row().cells
        for i, valor in enumerate(fila):
            celdas[i].text = ""
            set_run_font(celdas[i].paragraphs[0].add_run(str(valor)), "Calibri", 10)
    if anchos:
        for fila in t.rows:
            for i, ancho in enumerate(anchos):
                fila.cells[i].width = Cm(ancho)
    doc.add_paragraph()


def pie_de_pagina(section, texto):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(texto + "  |  Pag. "), "Calibri", 9, False, (0x5B, 0x5B, 0x5B))
    run = p.add_run()
    for tipo, texto_campo in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto_campo
        run._r.append(el)


def ejecutar_pruebas():
    """Corre un recorrido de pruebas contra la API y devuelve los resultados."""
    _, productores = pedir("GET", "/productores")
    _, tecnicos = pedir("GET", "/tecnicos")
    _, tipos = pedir("GET", "/tipos-cultivo")
    productor = next(p for p in productores if p["cuit"] == "30-12345678-9")
    tecnico = tecnicos[0]
    tipo = tipos[0]
    inexistente = "000000000000000000000000"

    resultados = []

    def prueba(metodo, ruta, cuerpo, descripcion, esperado):
        status, body = pedir(metodo, ruta, cuerpo)
        resultados.append({
            "metodo": metodo, "ruta": ruta, "cuerpo": cuerpo, "descripcion": descripcion,
            "esperado": esperado, "status": status, "body": body,
        })
        return body

    prueba("GET", "/api", None, "Listado de endpoints", 200)
    prueba("GET", "/productores", None, "Listar productores desde MongoDB", 200)
    prueba("GET", "/productores/" + productor["_id"], None, "Buscar productor por _id", 200)
    prueba("GET", "/productores/" + inexistente, None, "Productor inexistente", 404)
    prueba("GET", "/productores/abc", None, "Id con formato invalido (CastError)", 400)
    prueba("POST", "/productores", {"nombre": "Error", "cuit": "123"}, "CUIT con formato invalido (match)", 400)
    prueba("POST", "/productores", {"nombre": "Repetido", "cuit": productor["cuit"]}, "CUIT repetido (unique)", 400)
    prueba("GET", "/lotes?productor=" + productor["_id"], None, "Filtrar lotes por productor (populate)", 200)
    prueba("POST", "/lotes", {"nombre": "Lote Error", "hectareas": -5, "productor": productor["_id"], "tipoCultivo": tipo["_id"]},
           "Hectareas negativas (min)", 400)
    prueba("POST", "/lotes", {"nombre": "Lote Error", "hectareas": 5, "productor": inexistente, "tipoCultivo": tipo["_id"]},
           "Lote con productor inexistente", 400)
    nuevo = prueba("POST", "/lotes", {"nombre": "Lote Prueba", "hectareas": 25, "productor": productor["_id"], "tipoCultivo": tipo["_id"]},
                   "Alta de lote", 201)
    lote_id = nuevo["lote"]["_id"]
    prueba("PUT", "/lotes/" + lote_id, {"estado": "cosechado"}, "Actualizar estado del lote", 200)
    prueba("PUT", "/lotes/" + lote_id, {"estado": "inventado"}, "Estado fuera del enum", 400)
    obs = prueba("POST", "/observaciones", {"lote": lote_id, "tecnico": tecnico["_id"], "tipo": "plaga",
                                            "descripcion": "Presencia de chinche. Revisar en 5 dias"}, "Alta de observacion", 201)
    prueba("GET", "/observaciones?lote=" + lote_id, None, "Filtrar observaciones por lote", 200)
    prueba("DELETE", "/lotes/" + lote_id, None, "Baja de lote con observaciones", 400)
    prueba("DELETE", "/observaciones/" + obs["observacion"]["_id"], None, "Baja de observacion", 200)
    prueba("DELETE", "/lotes/" + lote_id, None, "Baja de lote sin observaciones", 200)
    prueba("DELETE", "/productores/" + productor["_id"], None, "Baja de productor con lotes asociados", 400)
    prueba("GET", "/no-existe", None, "Ruta inexistente", 404)
    return resultados


def escribir_html(resultados):
    filas = []
    for r in resultados:
        ok = "ok" if r["status"] == r["esperado"] else "fail"
        cuerpo = ""
        if r["cuerpo"] is not None:
            cuerpo = "<pre class='req'>" + json.dumps(r["cuerpo"], ensure_ascii=False, indent=2) + "</pre>"
        respuesta = json.dumps(r["body"], ensure_ascii=False, indent=2)
        if len(respuesta) > 700:
            respuesta = respuesta[:700] + "\n..."
        filas.append(f"""
        <article class="card">
          <div class="meta"><span class="method {r['metodo']}">{r['metodo']}</span><code>{r['ruta']}</code>
          <span class="status {ok}">{r['status']} (esperado {r['esperado']})</span></div>
          <h3>{r['descripcion']}</h3>{cuerpo}<pre>{respuesta}</pre>
        </article>""")
    total_ok = sum(1 for r in resultados if r["status"] == r["esperado"])
    html = f"""<!DOCTYPE html>
<html lang="es"><head><meta charset="utf-8"><title>Pruebas funcionales - Sistema agrario (MongoDB)</title>
<style>
body {{ font-family: Calibri, Arial, sans-serif; background: #f4f6f8; margin: 24px; color: #1b1b1b; }}
h1 {{ color: #1F4E79; margin-bottom: 4px; }}
.card {{ background: #fff; border: 1px solid #d0d7de; border-radius: 8px; padding: 12px 16px; margin-bottom: 12px; }}
.card h3 {{ margin: 6px 0; font-size: 15px; }}
.meta {{ display: flex; gap: 12px; align-items: center; }}
code {{ background: #eef2f6; padding: 3px 8px; border-radius: 4px; font-size: 13px; word-break: break-all; }}
pre {{ background: #0f172a; color: #e2e8f0; padding: 10px; border-radius: 6px; overflow: auto; font-size: 12px; margin: 6px 0 0; }}
pre.req {{ background: #eef2f6; color: #1b1b1b; }}
.method {{ color: #fff; padding: 3px 8px; border-radius: 4px; font-weight: bold; font-size: 13px; }}
.GET {{ background: #1d4ed8; }} .POST {{ background: #15803d; }} .PUT {{ background: #b45309; }} .DELETE {{ background: #b91c1c; }}
.status.ok {{ color: #15803d; font-weight: bold; }} .status.fail {{ color: #b91c1c; font-weight: bold; }}
</style></head><body>
<h1>Registro de pruebas funcionales</h1>
<p>Sistema agrario - monitoreo de lotes | Express + MongoDB (Mongoose) | {BASE} | {total_ok}/{len(resultados)} pruebas con el resultado esperado</p>
{''.join(filas)}
</body></html>
"""
    ruta = os.path.join(ROOT, "pruebas.html")
    with open(ruta, "w", encoding="utf-8") as f:
        f.write(html)
    return ruta


def main():
    resultados = ejecutar_pruebas()
    html_path = escribir_html(resultados)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(2)
    pie_de_pagina(section, NOMBRE + " - Parte 2 | Tecnicatura Superior en Desarrollo de Software - IFTS 29")

    if os.path.exists(LOGO):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(LOGO, width=Cm(3.2))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("DOCUMENTACION - Parte 2"), "Calibri", 22, True, AZUL)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("Desarrollo de Sistemas Web Backend - Segunda entrega (2do parcial)"), "Calibri", 12, False, (0x5B, 0x5B, 0x5B))

    parrafo(doc, "Caso: Sistema agricola / Monitoreo de lotes.")
    parrafo(doc, "Empresa de desarrollo: Carpincho (Grupo 9).")
    parrafo(doc, "Integrantes: Daniel Clementin, Rocio Espada, Cintia Wdaviak, Francisco Scuzzarello Lopez.")
    parrafo(doc, "Nombre del documento y de la carpeta Drive: " + NOMBRE)

    titulo(doc, "Indice")
    indice = [
        "1. Links", "2. Bibliografia", "3. Introduccion", "4. Integrantes, roles y responsabilidades",
        "5. Descripcion funcional de la aplicacion", "6. Modulos del proyecto", "7. Persistencia de datos con MongoDB",
        "8. Arquitectura y funcionamiento", "9. Validaciones y manejo de errores", "10. Requerimientos",
        "11. Instalacion y ejecucion", "12. Pruebas funcionales", "13. Uso de inteligencia artificial",
    ]
    for item in indice:
        parrafo(doc, item)

    titulo(doc, "1. Links")
    parrafo(doc, "Carpeta de Drive: " + DRIVE)
    parrafo(doc, "Repositorio GitHub: " + GITHUB)
    parrafo(doc, "Video explicativo (funcionamiento y uso de MongoDB): " + VIDEO)
    parrafo(doc, "Coleccion de pruebas: postman/Sistema-Agrario.postman_collection.json (dentro del repositorio).")

    titulo(doc, "2. Bibliografia")
    parrafo(doc, "Material de catedra:", True)
    vineta(doc, "Presentaciones y carpeta Code de las clases (Express, rutas, controladores, POO y JSON).")
    vineta(doc, "Consigna del 2do parcial IFTS 29 - 1C2026 y documento 'Segunda Entrega'.")
    parrafo(doc, "Documentacion oficial:", True)
    vineta(doc, "Mongoose - Getting Started. https://mongoosejs.com/docs/index.html")
    vineta(doc, "Mongoose - Validation. https://mongoosejs.com/docs/validation.html")
    vineta(doc, "Mongoose - Populate. https://mongoosejs.com/docs/populate.html")
    vineta(doc, "MongoDB Manual. https://www.mongodb.com/docs/manual/")
    vineta(doc, "Express - Manejo de errores. https://expressjs.com/es/guide/error-handling.html")
    vineta(doc, "Express - Uso de middleware. https://expressjs.com/es/guide/using-middleware.html")
    vineta(doc, "MDN - async function / await. https://developer.mozilla.org/es/docs/Web/JavaScript/Reference/Statements/async_function")
    vineta(doc, "dotenv. https://www.npmjs.com/package/dotenv")
    vineta(doc, "Pug. https://pugjs.org/api/getting-started.html")
    parrafo(doc, "Videos (YouTube): COMPLETAR si el grupo uso algun tutorial (titulo, canal y link).")

    titulo(doc, "3. Introduccion")
    parrafo(doc,
            "En esta segunda entrega se continuo el sistema de monitoreo de lotes de la primera entrega. "
            "El cambio principal es la persistencia: los archivos JSON se reemplazaron por una base de datos MongoDB, "
            "a la que se accede con Mongoose. Se mantuvieron los cinco modulos (productores, tipos de cultivo, lotes, "
            "tecnicos y observaciones), las rutas, las vistas Pug y la coleccion de pruebas.")
    parrafo(doc,
            "Ademas se incorporo programacion asincronica con async/await en todos los controladores, validaciones "
            "simples definidas en los Schemas, verificacion de las relaciones entre colecciones y un middleware "
            "centralizado de manejo de errores. Siguiendo la consigna, no se agregaron login, despliegue ni otras "
            "herramientas fuera de lo visto en clase, para evitar sumar complejidad innecesaria (overengineering).")
    parrafo(doc,
            "El caso plantea la falta de conectividad en el campo. Por eso el sistema esta pensado para que el tecnico "
            "registre las observaciones al volver a una zona con conexion, con la fecha real de la visita (el campo fecha "
            "se puede informar; si no se envia, se usa la fecha actual).")

    titulo(doc, "4. Integrantes, roles y responsabilidades")
    tabla(doc, ["Integrante", "Rol", "Responsabilidades en esta entrega"], anchos=[3.6, 4.2, 9.2], filas=[
        ["Daniel Clementin", "Desarrollador - Productores y Tecnicos",
         "Schemas Productor y Tecnico, controladores con async/await, validaciones (CUIT con formato y unico) y regla de baja."],
        ["Rocio Espada", "Desarrolladora - Lotes y Tipos de cultivo",
         "Schemas Lote y TipoCultivo con referencias (ObjectId + ref), populate, filtro por productor, enum de estado y min de hectareas."],
        ["Cintia Wdaviak", "Desarrolladora - Observaciones y vistas",
         "Schema Observacion, filtros por lote y tecnico, adaptacion de las vistas Pug y formularios al _id de MongoDB."],
        ["Francisco Scuzzarello Lopez", "Lider tecnico / documentacion y pruebas",
         "Conexion a MongoDB (config/db.js, .env), seed de datos, middleware de errores, coleccion de pruebas, repositorio, Drive y documentacion."],
    ])
    parrafo(doc, "Cada integrante explica en el video su parte del desarrollo.")

    titulo(doc, "5. Descripcion funcional de la aplicacion")
    titulo(doc, "5.1. Objetivo general", 13)
    parrafo(doc,
            "Permitir a una empresa agricola registrar productores, lotes, tipos de cultivo, tecnicos y observaciones "
            "de campo, consultando y actualizando esa informacion desde una API REST y desde vistas web simples, "
            "con los datos guardados en MongoDB.")
    titulo(doc, "5.2. Funcionalidades principales actuales", 13)
    for f in [
        "ABMC (alta, baja, modificacion y consulta) de los cinco modulos, persistido en MongoDB.",
        "Consulta por _id con respuesta 404 si no existe y 400 si el id no tiene formato valido.",
        "Relaciones entre colecciones con ObjectId y populate: el lote muestra su productor y cultivo; la observacion, su lote y tecnico.",
        "Filtros: lotes por productor y observaciones por lote y/o tecnico.",
        "Validaciones en los Schemas: obligatorios, minimos, valores permitidos (enum), formato de CUIT y valores unicos.",
        "Baja controlada: no se elimina un productor, cultivo, lote o tecnico que tenga datos asociados.",
        "Vistas Pug con formularios de alta; si hay un error se muestra una pantalla con el mensaje.",
        "Middleware: logger de peticiones, respuesta JSON o HTML segun quien llama, 404 de rutas y manejo de errores.",
        "Script de carga inicial (npm run seed) que pasa los datos de la primera entrega a MongoDB.",
    ]:
        vineta(doc, f)

    titulo(doc, "6. Modulos del proyecto")
    modulos = [
        ("Modulo de Productores", ["Listar, crear, editar y eliminar productores.", "CUIT obligatorio, con formato 00-00000000-0 y unico.",
                                   "No se borra un productor que tiene lotes."]),
        ("Modulo de Tipos de cultivo", ["Catalogo de cultivos (soja, maiz, trigo, etc.) con nombre unico.",
                                        "Ciclo en dias mayor a 0.", "No se borra un cultivo asignado a un lote."]),
        ("Modulo de Lotes", ["Cada lote referencia un productor y un tipo de cultivo existentes.",
                             "Hectareas mayores a 0; estado: activo, en barbecho o cosechado.",
                             "Filtro GET /lotes?productor=<id>.", "No se borra un lote con observaciones."]),
        ("Modulo de Tecnicos", ["Listar, crear, editar y eliminar tecnicos de campo.", "No se borra un tecnico que cargo observaciones."]),
        ("Modulo de Observaciones", ["Registro de humedad, suelo, plaga, clima u otro sobre un lote, hecho por un tecnico.",
                                     "Filtros GET /observaciones?lote=<id>&tecnico=<id>, ordenadas por fecha."]),
    ]
    for nombre, items in modulos:
        parrafo(doc, nombre + ":", True)
        for i in items:
            vineta(doc, i)
    parrafo(doc, "Proximos modulos (tercera etapa):", True)
    vineta(doc, "Login, roles y seguridad.")
    vineta(doc, "Despliegue y base de datos en la nube (MongoDB Atlas).")
    vineta(doc, "Modulo de inteligencia artificial sobre las observaciones cargadas.")

    titulo(doc, "7. Persistencia de datos con MongoDB")
    parrafo(doc,
            "La base se llama sistema_agrario. La cadena de conexion se lee de la variable MONGODB_URI del archivo .env, "
            "por lo que funciona igual con un MongoDB local o con Atlas. config/db.js se conecta con mongoose.connect "
            "usando async/await y el servidor arranca solo despues de conectarse.")
    tabla(doc, ["Coleccion", "Modelo", "Campos", "Referencias"], [
        ["productores", "Productor", "nombre*, cuit* (unico), localidad, telefono", "-"],
        ["tiposCultivo", "TipoCultivo", "nombre* (unico), descripcion, cicloDias (>0)", "-"],
        ["lotes", "Lote", "nombre*, hectareas* (>0), estado (enum)", "productor -> Productor, tipoCultivo -> TipoCultivo"],
        ["tecnicos", "Tecnico", "nombre*, apellido*, especialidad, matricula", "-"],
        ["observaciones", "Observacion", "fecha (default hoy), tipo* (enum), descripcion", "lote -> Lote, tecnico -> Tecnico"],
    ])
    parrafo(doc, "(*) campo obligatorio. Cada documento tiene un _id (ObjectId) generado por MongoDB.")
    parrafo(doc, "Ejemplo de Schema con referencias (models/Lote.js):")
    codigo(doc,
           'const loteSchema = new mongoose.Schema({\n'
           '    nombre: { type: String, required: [true, "nombre es obligatorio"], trim: true },\n'
           '    hectareas: { type: Number, required: true, min: [0.1, "hectareas debe ser mayor a 0"] },\n'
           '    productor: { type: mongoose.Schema.Types.ObjectId, ref: "Productor", required: true },\n'
           '    tipoCultivo: { type: mongoose.Schema.Types.ObjectId, ref: "TipoCultivo", required: true },\n'
           '    estado: { type: String, enum: ["activo", "en barbecho", "cosechado"], default: "activo" }\n'
           '});')
    parrafo(doc, "Operaciones de Mongoose utilizadas:")
    tabla(doc, ["Operacion", "Uso en el proyecto"], [
        ["find / findById", "Listados, filtros y consulta por id"],
        ["create", "Altas"],
        ["findByIdAndUpdate (runValidators)", "Modificaciones validando el Schema"],
        ["findByIdAndDelete", "Bajas"],
        ["exists", "Verificar referencias y datos asociados antes de borrar"],
        ["populate", "Traer los datos del documento referenciado (productor, cultivo, lote, tecnico)"],
        ["sort", "Ordenar listados por nombre o fecha"],
    ])
    parrafo(doc,
            "Los archivos de data/ de la primera entrega se conservaron como datos de ejemplo: seed.js los lee, vacia "
            "las colecciones y los inserta en MongoDB, reemplazando los id numericos por los ObjectId generados.")

    titulo(doc, "8. Arquitectura y funcionamiento")
    parrafo(doc, "Se mantiene la estructura modular por capas de la primera entrega:")
    codigo(doc,
           "index.js        servidor Express, middlewares y montaje de rutas\n"
           "config/db.js    conexion a MongoDB\n"
           "models/         Schemas y modelos de Mongoose\n"
           "routes/         rutas HTTP de cada modulo\n"
           "controllers/    logica de cada modulo con async/await\n"
           "middleware/     logger, responder y errorHandler\n"
           "views/          plantillas Pug\n"
           "data/ + seed.js datos de ejemplo y carga inicial")
    parrafo(doc, "Recorrido de una peticion (ejemplo POST /lotes):")
    for paso in [
        "logger muestra en consola el metodo y la ruta; express.json / urlencoded leen el cuerpo.",
        "routes/lotesRoutes.js deriva la peticion a crearLote del controlador.",
        "El controlador verifica con await que existan el productor y el tipo de cultivo.",
        "Lote.create valida el Schema y guarda el documento en MongoDB.",
        "responder devuelve JSON 201 a Postman / Thunder Client o redirige a /ver/lotes si vino del formulario.",
        "Si algo falla, el catch llama a next(error) y errorHandler arma la respuesta 400 / 500.",
    ]:
        vineta(doc, paso)

    titulo(doc, "9. Validaciones y manejo de errores")
    tabla(doc, ["Situacion", "Donde se controla", "Respuesta"], [
        ["Campo obligatorio faltante, valor fuera de rango o de enum, CUIT mal formado", "Schema de Mongoose (ValidationError)", "400 con el mensaje del Schema"],
        ["CUIT o nombre de cultivo repetido", "Indice unique (error 11000)", "400 'Ya existe un registro con ese ...'"],
        ["Id con formato invalido (ej: /productores/abc)", "CastError de Mongoose", "400 'El valor de _id no es valido'"],
        ["Referencia inexistente (productor, cultivo, lote, tecnico)", "Controlador (exists)", "400"],
        ["Baja de un registro con datos asociados", "Controlador (exists)", "400"],
        ["Registro inexistente", "Controlador", "404"],
        ["Ruta inexistente", "middleware noEncontrado", "404"],
        ["Error inesperado", "middleware manejarErrores", "500"],
    ])
    parrafo(doc,
            "Como Express 4 no captura solo los errores de funciones async, cada controlador usa try/catch y pasa el "
            "error con next(error) al middleware de errores, que se registra al final de index.js.")

    titulo(doc, "10. Requerimientos")
    parrafo(doc, "Funcionales:", True)
    for rf in [
        "RF01 - Registrar, consultar, modificar y eliminar productores, tecnicos, tipos de cultivo, lotes y observaciones.",
        "RF02 - Asociar cada lote a un productor y un tipo de cultivo existentes.",
        "RF03 - Registrar observaciones de un tecnico sobre un lote, con fecha y tipo.",
        "RF04 - Filtrar lotes por productor y observaciones por lote o tecnico.",
        "RF05 - Impedir la baja de registros con datos asociados.",
        "RF06 - Cargar datos desde formularios web simples.",
    ]:
        vineta(doc, rf)
    parrafo(doc, "No funcionales:", True)
    for rnf in [
        "RNF01 - Persistencia en MongoDB mediante Mongoose.",
        "RNF02 - Arquitectura modular (rutas, controladores, modelos, middleware, vistas).",
        "RNF03 - Operaciones de base de datos asincronicas con async/await.",
        "RNF04 - Respuestas JSON con codigos HTTP adecuados (200, 201, 400, 404, 500).",
        "RNF05 - Configuracion por variables de entorno (.env), sin credenciales en el repositorio.",
    ]:
        vineta(doc, rnf)

    titulo(doc, "11. Instalacion y ejecucion")
    parrafo(doc, "Requisitos: Node.js 18+ y MongoDB (local, Docker o Atlas).")
    codigo(doc,
           "git clone " + GITHUB + "\n"
           "cd sistema-agrario\n"
           "npm install\n"
           "copy .env.example .env     (en Linux / Mac: cp .env.example .env)\n"
           "npm run seed               (carga los datos de ejemplo)\n"
           "npm start                  (http://localhost:3000)")
    parrafo(doc, "Para levantar MongoDB con Docker: docker run -d --name mongo-agrario -p 27017:27017 mongo:7")

    titulo(doc, "12. Pruebas funcionales")
    total_ok = sum(1 for r in resultados if r["status"] == r["esperado"])
    parrafo(doc,
            "Las pruebas se hicieron sobre http://localhost:3000 con la base cargada por npm run seed. "
            "La coleccion postman/Sistema-Agrario.postman_collection.json (34 peticiones, se importa en Postman o "
            "Thunder Client) guarda los _id en variables y se puede correr completa en orden. "
            f"Resumen de la ejecucion: {total_ok} de {len(resultados)} pruebas devolvieron el codigo esperado.")
    tabla(doc, ["Metodo", "Ruta", "Prueba", "Esp.", "Real"], [
        [r["metodo"], r["ruta"] if len(r["ruta"]) < 45 else r["ruta"][:42] + "...", r["descripcion"], r["esperado"], r["status"]]
        for r in resultados
    ], anchos=[1.8, 7.0, 5.6, 1.3, 1.3])
    parrafo(doc, "Detalle de algunas respuestas:")
    for r in resultados:
        if r["descripcion"] in ("Alta de lote", "CUIT con formato invalido (match)", "Id con formato invalido (CastError)",
                                "Baja de productor con lotes asociados", "Estado fuera del enum"):
            parrafo(doc, f"{r['metodo']} {r['ruta']}  ->  {r['status']}  ({r['descripcion']})", True)
            if r["cuerpo"] is not None:
                codigo(doc, "Body: " + json.dumps(r["cuerpo"], ensure_ascii=False))
            codigo(doc, json.dumps(r["body"], ensure_ascii=False, indent=2))
    parrafo(doc, "Capturas:", True)
    for leyenda, imagen in CAPTURAS:
        if os.path.exists(imagen):
            doc.add_picture(imagen, width=Cm(16.5))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run(leyenda), "Calibri", 9, False, (0x5B, 0x5B, 0x5B))

    titulo(doc, "13. Uso de inteligencia artificial")
    parrafo(doc,
            "Se utilizo asistencia de IA para planificar la migracion de JSON a MongoDB, revisar que el alcance coincida "
            "con la consigna (sin agregar herramientas no pedidas) y generar la coleccion de pruebas y este documento. "
            "El codigo fue revisado y probado por el grupo. Prompts usados:")
    vineta(doc, "Migrar el CRUD de monitoreo de lotes de archivos JSON a MongoDB con Mongoose, usando async/await, "
                "middleware de errores y validaciones simples, manteniendo la estructura de la primera entrega.")
    vineta(doc, "Armar la documentacion de la segunda entrega siguiendo la plantilla de la catedra.")

    doc.save(SALIDA)
    print("HTML:", html_path)
    print("DOCX:", SALIDA)


if __name__ == "__main__":
    main()
